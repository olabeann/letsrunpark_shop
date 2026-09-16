from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서.docx")
NAVY, PALE, GRID = "1F4E79", "D9EAF7", "D9D9D9"

def font(run, size=10, bold=False, color="000000"):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size, run.font.bold = Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)

def border(cell):
    pr = cell._tc.get_or_add_tcPr()
    b = pr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders"); pr.append(b)
    for side in ("top", "left", "bottom", "right"):
        x = b.find(qn("w:" + side))
        if x is None:
            x = OxmlElement("w:" + side); b.append(x)
        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "6"); x.set(qn("w:color"), GRID)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); x = OxmlElement("w:shd")
    x.set(qn("w:fill"), fill); pr.append(x)

def cell(c, value, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color="000000", size=9):
    c.text = ""; p = c.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(value)); font(r, size, bold, color)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(c)

def table(doc, headers, rows, widths, money_cols=()):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, h in enumerate(headers):
        t.columns[i].width = Cm(widths[i]); cell(t.rows[0].cells[i], h, True, WD_ALIGN_PARAGRAPH.CENTER, "FFFFFF"); shade(t.rows[0].cells[i], NAVY)
    for ri, row in enumerate(rows):
        cs = t.add_row().cells
        for i, value in enumerate(row):
            cell(cs[i], value, False, WD_ALIGN_PARAGRAPH.RIGHT if i in money_cols else WD_ALIGN_PARAGRAPH.LEFT)
            if ri % 2: shade(cs[i], "F5F7F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(11); p.paragraph_format.space_after = Pt(5)
    font(p.add_run(text), 13, True)

def body(doc, text, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), 9.5)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document(); sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2); sec.right_margin = Cm(2)
doc.styles["Normal"].font.name = "Malgun Gothic"; doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8); font(p.add_run("견 적 서"), 22, True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(15); font(p.add_run("말마프렌즈 온라인 스토어 추가 개발"), 13, True, NAVY)

t = table(doc, [], [], [])
t._element.getparent().remove(t._element)
meta = doc.add_table(rows=3, cols=4); meta.alignment = WD_TABLE_ALIGNMENT.CENTER
for r, values in enumerate([
    ("수신", "렛츠런파크 담당부서 귀중", "견적일", "2026. 09. 16."),
    ("공급자", "주식회사 위이", "견적 유효기간", "별도 협의"),
    ("사업자번호", "236-81-02950", "수행기간", "착수일로부터 25~30영업일(예정)"),
]):
    for c, value in enumerate(values):
        cell(meta.rows[r].cells[c], value, c in (0, 2), WD_ALIGN_PARAGRAPH.CENTER if c in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT)
        if c in (0, 2): shade(meta.rows[r].cells[c], PALE)
doc.add_paragraph().paragraph_format.space_after = Pt(2)
table(doc, ["구분", "금액 및 조건"], [
    ("시스템 추가 구축비 (VAT 포함)", "금 팔백육십육만팔천원정  (₩ 8,668,000)"),
    ("추가 인프라 운영비", "기존 인프라 용량 내 운영 시 별도 증액 없음"),
    ("초년도 추가 소요 비용 (VAT 포함)", "금 팔백육십육만팔천원정  (₩ 8,668,000)"),
], [6.2, 9.5], {1})

heading(doc, "1. 견적 산정 기준")
body(doc, "본 견적은 기존 렛츠런플레이 통합예약결제시스템의 회원, 로그인, 관리자 인증, 결제 모듈 및 서버 인프라를 재사용하는 것을 전제로 합니다. 관리자 시스템을 별도 구축하지 않고 기존 예약 운영 관리자에 굿즈 운영 메뉴를 추가하며, 사용자용 온라인 스토어 화면과 굿즈 주문 도메인을 신규 개발합니다.")
body(doc, "기획·디자인 및 QA는 1 PD당 160,000원, 개발 및 인프라 작업은 1 PD당 220,000원을 적용하였습니다. 부가가치세는 별도 산정 후 합계에 반영합니다.")

heading(doc, "2. 구축비 산정 내역")
table(doc, ["구분", "수량", "일 단가(원)", "금액(원)"], [
    ("기획 및 기능명세·화면 설계·QA", "8 PD", "160,000", "1,280,000"),
    ("개발 - 상품·주문·재고·배송 도메인 및 사용자 스토어", "18 PD", "220,000", "3,960,000"),
    ("개발 - 기존 관리자 메뉴 확장·출고 엑셀·통합 테스트", "11 PD", "220,000", "2,420,000"),
    ("인프라·배포 환경 보완", "1 PD", "220,000", "220,000"),
    ("공급가액 (VAT 별도)", "", "", "7,880,000"),
    ("부가가치세 (10%)", "", "", "788,000"),
    ("합계 금액 (VAT 포함)", "", "", "8,668,000"),
], [6.8, 2.3, 3.2, 3.4], {2, 3})

heading(doc, "3. 주요 구축 범위")
table(doc, ["구분", "포함 범위"], [
    ("사용자 온라인 스토어", "상품 목록, 상품 상세, 상품 이미지·설명·가격·판매 상태 표시, 수량 선택"),
    ("장바구니 및 주문", "복수 상품 장바구니, 수량 변경·삭제, 배송비·총 결제금액 계산, 배송지 입력"),
    ("결제 및 주문 조회", "기존 회원 로그인·PG 결제 재사용, 주문번호 발급, 주문 조회 및 출고 전 전체 취소"),
    ("상품·재고 관리", "상품 등록·수정, 대표 이미지 1장, 가격·재고·판매 상태·노출 순서 관리, 재고 이력"),
    ("굿즈 판매 현황", "주문·배송지 조회, 출고 대기 주문 엑셀 다운로드, 출고 완료 일괄 처리, 송장번호 수동 입력"),
    ("재고 처리", "결제 직전 재고 재확인, 결제 성공 시 차감, 전체 취소 완료 시 주문상품 재고 복원"),
], [4.0, 11.7])

doc.add_page_break()
heading(doc, "4. 관리자 운영 방식")
body(doc, "관리자 페이지는 새로 구축하지 않습니다. 기존 예약 운영 관리자에 아래 메뉴를 추가하여 기존 권한 및 운영 체계 안에서 관리합니다.")
table(doc, ["추가 메뉴", "주요 기능"], [
    ("굿즈 판매 현황", "주문번호·주문일·구매자·상품·수량·결제금액 조회, 주문 상세·배송지 확인, 출고 대기 엑셀 다운로드, 출고 완료 일괄 처리, 출고 전 전체 취소"),
    ("온라인 스토어 운영", "상품 등록·수정, 대표 이미지 등록, 상품명·설명·판매가격·재고 입력, 판매 중·품절·숨김 상태 변경, 사용자 화면 노출 순서 설정"),
], [4.0, 11.7])

heading(doc, "5. 주문 및 출고 처리 기준")
table(doc, ["단계", "처리 기준"], [
    ("장바구니", "상품을 담는 시점에는 재고를 선점하거나 차감하지 않습니다."),
    ("결제 요청", "결제 직전에 최신 재고를 확인하며, 재고 부족 시 결제를 중단하고 안내합니다."),
    ("결제 성공", "주문과 주문상품을 저장하고 상품별 재고를 차감합니다."),
    ("출고 운영", "관리자가 주 1회 출고 대기 주문을 엑셀로 다운로드하고 포장·택배 접수 후 출고 완료로 일괄 처리합니다."),
    ("전체 취소", "출고 전 주문에 한해 주문 전체를 취소할 수 있으며, PG 취소 완료 후 모든 주문상품 재고를 복원합니다."),
], [3.0, 12.7])

heading(doc, "6. 일정 계획")
table(doc, ["구분", "공수", "기간", "주요 작업"], [
    ("정책 확정·화면 설계", "8 PD", "1주차", "배송비·출고·취소·재고 정책 확정, 기능명세 및 화면 설계"),
    ("핵심 개발", "20 PD", "1~4주차", "상품·주문·재고·배송 데이터, 사용자 스토어, 결제·전체 취소 연동"),
    ("관리자 확장·검수", "10 PD", "3~5주차", "기존 관리자 메뉴 추가, 출고 엑셀, 통합 테스트, 배포 및 인수"),
], [3.2, 2.2, 3.0, 7.3])

heading(doc, "7. 제외 및 별도 협의 항목")
for x in [
    "체험 예약권과 굿즈의 혼합 장바구니 및 혼합 결제",
    "상품 옵션 조합, 쿠폰·포인트·회원등급 할인, 부분취소·부분환불, 교환·반품 자동 접수",
    "택배사 API 연동, 자동 송장 출력, 자동 배송조회, 주문 상태 알림톡·문자",
    "상품 후기·문의, 찜하기·추천 상품, 통계 대시보드, 비회원 구매 및 해외 배송",
    "PG 수수료, 카드 수수료, 메시지 발송 비용, 추가 이미지 제작 및 보안성 검토·인증 대응",
]: body(doc, x, True)

heading(doc, "8. 전제 및 유의 사항")
for x in [
    "기존 결제 모듈이 상품 주문의 결제 요청 및 전체 취소 처리를 지원하거나 이에 준하는 연동이 가능해야 합니다.",
    "기존 서버·DB·스토리지 용량 내 운영을 전제로 하며, 이미지 저장량·트래픽·보안 요건에 따른 인프라 증설은 AWS 실비를 별도 협의합니다.",
    "상품 수, 배송비 정책, 출고 요일 및 마감 시각, 구매 제한 수량, 최초 상품 이미지 등 운영 정책과 자료는 착수 전 제공되어야 합니다.",
    "확정 기능명세 외 신규 기능, 정책 변경, 디자인 구조 변경 및 운영 지원 요청은 별도 공수로 산정합니다.",
]: body(doc, x, True)

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("주식회사 위이  |  말마프렌즈 온라인 스토어 추가 개발 견적서"), 8, False, "666666")
doc.save(OUT); print(OUT)
