from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
SRC=ROOT/'planning-v2.1'
OUT=ROOT/'planning-v2.2';OUT.mkdir(exist_ok=True)
fs=json.loads((SRC/'기능목록.json').read_text())
old=Document(SRC/'03_말마프렌즈_개발전달서_v2.1.docx')
d=Document(SRC/'03_말마프렌즈_개발전달서_v2.1.docx')
for e in list(d.element.body):
 if e.tag!=qn('w:sectPr'):d.element.body.remove(e)
md=[]
def p(text,style=None):
 para=d.add_paragraph(text,style);md.append(text+'\n');return para
def h(text):md.append('## '+text+'\n');return d.add_heading(text,1)
def page():d.add_page_break()
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,row in enumerate([headers]+rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,(cell,s) in enumerate(zip(cells,row)):
   cell.width=Cm(widths[j]);cell.text=s
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173C55' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
    for run in para.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
  cant=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(cant)
  if i==0:
   rep=OxmlElement('w:tblHeader');t.rows[i]._tr.get_or_add_trPr().append(rep)
 md.append('| '+' | '.join(headers)+' |\n|'+ '|'.join(['---']*len(headers))+'|\n'+'\n'.join('| '+' | '.join(r).replace('\n','<br>')+' |' for r in rows)+'\n')
 return t
summaries=[
'기존 서비스에서 스토어 목록으로 진입','판매 상태와 지정 순서에 맞춰 상품 노출','상세 콘텐츠와 구매 및 배송 취소 안내 표시','재고 범위에서 수량 선택 및 단일 상품 바로 구매','기존 회원 인증 후 원래 구매 행동 복원','복수 상품 수량 삭제 및 주문 합계 계산','주문자와 수령인을 구분하고 배송지 입력','주소 검색 또는 직접 입력 및 필수 안내 동의','서버 금액 재검증 후 카드 승인 주문 확정 재고 차감','확정 주문번호 내역 금액과 주문 조회 제공','로그인 회원의 주문만 최신순 조회','주문 당시 금액 정책 수령 정보 보존 표시','허용 기한 내 발송 전 주문의 전액 취소','수동 등록 송장 표시 및 외부 배송조회',
'기간 상태 키워드 조건으로 전체 주문 조회','검색 결과를 상품별 행의 CSV로 내려받기','권한 범위에서 주문 배송 결제 정보 조회','허용 주문 전액 취소 및 사유 이력 저장','대기 주문 선택 및 직접 입력 송장 검증','선택 주문을 재검증하고 발송 완료 저장','판매 상태와 상품명으로 상품 목록 검색','상품명 가격 재고 요약과 대표 이미지 저장','서식 이미지 링크 영상으로 상세 콘텐츠 편집','재고와 판매 상태 변경 및 변경 이력 보존','지정 순서를 사용자 목록에 반영','배송비 취소시간 배송 안내를 새 주문부터 적용']
flags={3:'상세 에디터 범위 확인',8:'동의 문구 확정',16:'CSV 유지 송장 업로드 제외',19:'공통 택배사 UI 제안',22:'파일 저장 연동 필요',23:'추가 범위 합의',25:'시연 미구현',26:'추가 범위 합의'}
def ids(f):return ', '.join(dict.fromkeys(re.findall(r'DEC-\d{3}',str(f)))) or '—'
def scope(n,f):return flags.get(n,'기존 연동 확인' if n in [5,9] else ('정책 확정 필요' if ids(f)!='—' else '기본 개발 범위'))
def link(para,text,anchor):
 para.clear();el=OxmlElement('w:hyperlink');el.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'173C55');pr.append(c);r.append(pr);txt=OxmlElement('w:t');txt.text=text;r.append(txt);el.append(r);para._p.append(el)
p('말마프렌즈 온라인 스토어 개발 전달서','Title')
p('전체 기능과 구현 기준을 함께 확인하는 개발 문서')
p('기획 개정안 v2.2  |  2026년 9월 18일')
p('회원의 굿즈 구매부터 관리자 출고와 상품 운영까지 총 26개 기능을 이 문서에 모았다. 기능표에서 전체 개발 범위를 확인하고 기능명을 클릭해 상세 동작으로 이동한다. 뒤쪽에는 공통 정책과 권한, 주문 상태, 데이터와 연동, 미확정 결정 항목, 검수 시나리오를 수록했다.')
p('관리자는 하나의 프로그램 안에서 예약관리시스템과 커머스관리시스템으로 나뉜다. 이번 개발 대상은 커머스 영역이다. 실제 회원 인증과 PG 및 서버 저장 연동은 구현해야 하며 시연 화면의 동작만으로 완료를 판단하지 않는다.')
h('사용자 기능 전체 목록')
rows=[]
for n,f in enumerate(fs[:14],1):rows.append([f['id']+'\n'+f['name'],f['screen'],summaries[n-1],scope(n,f)])
t=table(['기능','화면','핵심 동작','개발 시 확인'],rows,[3.8,2.5,6.8,4.2])
for i,f in enumerate(fs[:14],1):link(t.rows[i].cells[0].paragraphs[0],f['id']+' '+f['name'],f['id'].replace('-','_'))
page();h('관리자 기능 전체 목록')
p('커머스관리시스템의 굿즈 판매 현황은 주문 조회와 출고를, 온라인 스토어 운영은 상품과 정책을 관리한다. 송장 엑셀 업로드와 양식 다운로드는 제외하며 목록 CSV 내려받기는 유지한다.')
rows=[]
for n,f in enumerate(fs[14:],15):rows.append([f['id']+'\n'+f['name'],f['screen'],summaries[n-1],scope(n,f)])
t=table(['기능','화면','핵심 동작','개발 시 확인'],rows,[3.8,2.5,6.8,4.2])
for i,f in enumerate(fs[14:],1):link(t.rows[i].cells[0].paragraphs[0],f['id']+' '+f['name'],f['id'].replace('-','_'))
h('공통 구현 기준')
p('회원과 관리자 권한은 서버에서 검사한다. 금액과 재고는 결제 직전에 재검증하고 주문 당시 값을 보존한다. 승인 취소 발송 요청은 중복 실행과 동시 요청에도 한 번만 반영한다. 로딩 오류 빈 결과를 구분하고 입력 오류는 해당 위치에 표시한다.')
p('관리자 표의 제목과 내용은 열 가운데에 정렬한다. 드롭다운 화살표는 오른쪽 세로 중앙에 배치한다. 표 아래 작업 완료 문구는 제거하고 변경된 상태로 결과를 확인한다. 송장 입력의 공통 택배사 중심 UI는 추가 반영 제안이다.')
p('배송비 3,000원과 취소 24시간은 확정 정책이 아니다. DEC 항목을 확인한 뒤 개발한다. 부분 취소 분할 배송 자동 반품 택배 API 자동 알림 비회원 구매는 이번 범위에서 제외한다.')
for start in range(0,26,3):
 page();h('기능별 개발 동작 '+str(start//3+1))
 for n,f in enumerate(fs[start:start+3],start+1):
  para=d.add_heading(f['id']+' '+f['name'],2);md.append('### '+f['id']+' '+f['name']+'\n')
  b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(n));b.set(qn('w:name'),f['id'].replace('-','_'));para._p.insert(0,b);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(n));para._p.append(e)
  for label,value in [('연결',f['screen']+' | '+f['basis']),('시작',f['trigger']),('개발 동작',f['success']),('검증과 예외',f['validation']),('현재 시연',f['current']),('검수',f['accept'])]:
   para=p(label+'  '+value);para.paragraph_format.space_after=Pt(5);para.paragraph_format.line_spacing=1.08
   for r in para.runs:r.font.size=Pt(9)
page()
# Retain all shared policy and handoff details, starting with the service scope heading.
started=False
for e in old.element.body:
 if e.tag==qn('w:sectPr'):continue
 text=''.join(e.itertext())
 if '1 서비스 소개와 개발자가 알아야 할 범위' in text:started=True
 if started:d.element.body.insert(len(d.element.body)-1,deepcopy(e))
for para in d.paragraphs:
 if para.text.startswith('IA에서 화면과 이동을 이해한 뒤'):
  para.text='앞쪽 전체 기능표와 기능별 개발 동작에서 FN-001부터 FN-026까지 확인한다. 이 문서 2장부터 5장은 공통 정책 데이터 연동 기준이다. 6장의 DEC 결정 항목을 해소하고 7장의 TC 검수 시나리오로 완료 여부를 판단한다. IA와 기능명세서는 화면 구조와 상세 요구를 대조할 때 참고한다.'
for sec in d.sections:
 for para in sec.footer.paragraphs:
  for r in para.runs:r.text=r.text.replace('v2.1.1','v2.2').replace('v2.1','v2.2')
path=OUT/'03_말마프렌즈_개발전달서_v2.2.docx';d.save(path)
base=(SRC/'03_개발전달서.md').read_text();base=base[base.index('## 1 서비스 소개'):].replace('v2.1','v2.2')
base=base.replace('IA에서 화면과 이동을 이해한 뒤 기능명세서 FN-001부터 FN-026의 동작을 확인한다.','앞쪽 전체 기능표와 기능별 개발 동작에서 FN-001부터 FN-026까지 확인한다.')
(OUT/'03_개발전달서.md').write_text('# 말마프렌즈 온라인 스토어 개발 전달서 v2.2\n\n'+'\n'.join(md)+'\n'+base)
print(path)
