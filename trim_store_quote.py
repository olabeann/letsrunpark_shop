from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.oxml.ns import qn

SOURCE = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서_v7.docx")
TARGET = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서_v8.docx")


def replace_paragraph(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def replace_cell(cell, value):
    replace_paragraph(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        replace_paragraph(paragraph, "")


copyfile(SOURCE, TARGET)
doc = Document(TARGET)

# 금액 요약에서 '0원'으로 보일 수 있는 표현을 제거하고 기존 운영계약 적용으로 명확히 정리합니다.
replace_cell(doc.tables[1].rows[1].cells[0], "인프라 운영")
replace_cell(doc.tables[1].rows[1].cells[1], "기존 통합예약결제시스템 운영계약 적용")
replace_cell(doc.tables[1].rows[2].cells[0], "추가 개발비 합계 (VAT 포함)")

# 고객사 제출용 견적서에 필요한 범위까지만 유지합니다.
# 표 7은 '전제 및 유의 사항'이며, 이후의 인프라 사양·0원 운영비·차년도 비용 표를 모두 삭제합니다.
last_keep = doc.tables[7]._element
body = doc._element.body
node = last_keep.getnext()
while node is not None and node.tag != qn("w:sectPr"):
    next_node = node.getnext()
    body.remove(node)
    node = next_node

doc.save(TARGET)
print(TARGET)
