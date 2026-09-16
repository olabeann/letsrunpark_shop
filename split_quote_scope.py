from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서_v8.docx")
TARGET = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서_v9.docx")


def replace_paragraph(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def replace_cell(cell, value):
    replace_paragraph(cell.paragraphs[0], value)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)


def set_row(table, index, values):
    for cell, value in zip(table.rows[index].cells, values):
        replace_cell(cell, value)


def shade(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_group_cell(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
    shade(cell)


copyfile(SOURCE, TARGET)
doc = Document(TARGET)

# 3. 주요 구축 범위: 사용자와 관리자를 한눈에 구분합니다.
scope = doc.tables[3]
for row, values in enumerate([
    ["구분", "포함 범위"],
    ["사용자 페이지", "스토어 메인: 판매 상품 목록, 대표 이미지, 상품명, 가격, 판매 중·품절 상태 표시"],
    ["", "상품 상세·장바구니: 상품 설명·재고 확인, 수량 선택, 복수 상품 담기, 수량 변경·삭제"],
    ["", "주문·결제·조회: 기존 회원 로그인 및 PG 결제, 배송지 입력, 주문번호 발급, 주문 조회, 출고 전 전체 취소"],
    ["관리자 기능", "굿즈 판매 현황: 주문·구매자·상품·결제금액·배송지 조회, 출고 대상 엑셀, 출고 완료 일괄 처리, 전체 취소"],
    ["", "온라인 스토어 운영: 상품·대표 이미지·가격·재고 등록 및 수정, 판매 상태와 노출 순서 관리"],
]):
    set_row(scope, row, values)

user_group = scope.cell(1, 0).merge(scope.cell(3, 0))
admin_group = scope.cell(4, 0).merge(scope.cell(5, 0))
replace_cell(user_group, "사용자 페이지")
replace_cell(admin_group, "관리자 기능")
format_group_cell(user_group)
format_group_cell(admin_group)

# 4. 기능 범위: 상세 기능명에도 담당 화면을 표시합니다.
detail = doc.tables[4]
labels = [
    "기능 영역",
    "사용자 - 화면·디자인",
    "사용자 - 회원·로그인",
    "사용자 - 상품·장바구니",
    "공통 - 재고 처리",
    "사용자 - 결제·전체 취소",
    "사용자 - 배송정보",
    "사용자 - 주문 조회",
    "관리자 - 굿즈 판매 현황",
    "관리자 - 온라인 스토어 운영",
    "제외 기능",
]
for row, label in enumerate(labels):
    replace_cell(detail.rows[row].cells[0], label)

doc.save(TARGET)
print(TARGET)
