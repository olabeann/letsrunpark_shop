from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_기능명세서_v1.docx")
REFERENCE = Path(r"C:\GitHub\letsrunpark_shop\outputs\말마프렌즈_온라인스토어_추가개발_견적서_v9.docx")
NAVY, LIGHT, GRID, ALT = "1F4E79", "D9EAF7", "D9D9D9", "F5F7F9"


def set_font(run, size=10, bold=False, color="000000"):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        node = borders.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)


def margins(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell(cell, value, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color="000000", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(str(value)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    border(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_table(doc, headers, rows, widths, centers=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        table.columns[i].width = Cm(widths[i])
        set_cell(table.rows[0].cells[i], header, True, WD_ALIGN_PARAGRAPH.CENTER, "FFFFFF", 8.5)
        shade(table.rows[0].cells[i], NAVY)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in centers else WD_ALIGN_PARAGRAPH.LEFT
            set_cell(cells[i], value, False, align)
            if row_index % 2:
                shade(cells[i], ALT)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), 13 if level == 1 else 11, True)


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), 9.5, True)
        set_font(p.add_run(text[len(bold_lead):]), 9.5)
    else:
        set_font(p.add_run(text), 9.5)


doc = Document()
ref = Document(REFERENCE)
src = ref.sections[0]
sec = doc.sections[0]
sec.page_width, sec.page_height = src.page_width, src.page_height
sec.top_margin, sec.bottom_margin = src.top_margin, src.bottom_margin
sec.left_margin, sec.right_margin = src.left_margin, src.right_margin
sec.header_distance, sec.footer_distance = src.header_distance, src.footer_distance

normal = doc.styles["Normal"]
normal.font.name = "Malgun Gothic"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
normal.font.size = Pt(9.5)

p = doc.add_paragraph()
p.style = doc.styles["Title"]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("말마프렌즈 온라인 스토어 기능명세서"), 20, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
set_font(p.add_run("MVP 개발 범위 및 운영 정책 정의"), 12, True, NAVY)

add_table(doc, ["문서 정보", "내용", "문서 정보", "내용"], [
    ("문서 버전", "v1.0", "작성일", "2026. 09. 16."),
    ("대상 시스템", "렛츠런플레이", "작성 주체", "주식회사 위이"),
    ("개발 형태", "사용자 스토어 신규 구축", "관리자", "기존 관리자 메뉴 확장"),
], [3.0, 5.0, 3.0, 5.0], {0, 2})

heading(doc, "1. 문서 목적")
body(doc, "본 문서는 한국마사회가 요청한 말마프렌즈 온라인 스토어의 MVP 개발 범위와 기능 동작 기준을 정의합니다. 사용자 스토어는 신규 구축하고, 관리자 기능은 기존 예약 운영 관리자에 메뉴를 추가하는 방식으로 구성합니다.")
body(doc, "기능은 고객사 메일에 명시된 요구사항, 온라인 판매를 위해 필요한 기본 기능, 착수 전 협의가 필요한 선택 기능으로 구분합니다. 본 문서에서 협의로 표시한 항목은 운영 정책 확정 후 최종 범위에 반영합니다.")

heading(doc, "2. 요구사항 구분")
add_table(doc, ["구분", "의미", "계약 적용 기준"], [
    ("필수", "고객사 요청에 직접 명시된 기능", "MVP 기본 범위에 포함"),
    ("운영 필수", "요청 기능을 실제로 운영하기 위해 필요한 기본 기능", "MVP 기본 범위에 포함"),
    ("협의", "정책 또는 상세 방식이 확정되지 않은 기능", "착수 전 확정하며 변경 시 공수 협의"),
    ("제외", "MVP에서 제공하지 않는 기능", "추가 요청 시 별도 견적"),
], [3.0, 7.0, 6.0], {0})

heading(doc, "3. 시스템 구성 및 범위 경계")
add_table(doc, ["영역", "구축 방식", "범위"], [
    ("사용자 온라인 스토어", "신규 구축", "상품 목록·상세, 장바구니, 주문·결제, 주문 조회, 전체 취소"),
    ("굿즈 판매 현황", "기존 관리자 메뉴 추가", "주문·배송지 조회, 출고 대상 관리, 출고 완료, 전체 취소"),
    ("온라인 스토어 운영", "기존 관리자 메뉴 추가", "상품·가격·이미지·재고·판매 상태·노출 순서 관리"),
    ("공통 시스템", "기존 기능 재사용", "회원 로그인, 관리자 인증·권한, PG 결제, 서버·DB 인프라"),
    ("상품 소개 콘텐츠", "고객사 제작", "고객사가 제작하는 인형별 정보제공 페이지에서 온라인 스토어로 연결"),
], [4.0, 4.0, 8.0], {0, 1})

doc.add_page_break()
heading(doc, "4. 사용자 페이지 기능")
user_rows = [
    ("USR-01", "필수", "스토어 진입", "렛츠런플레이 메뉴 또는 고객사가 제작한 상품 소개 페이지에서 온라인 스토어로 이동한다.", "소개 페이지 제작은 제외"),
    ("USR-02", "운영 필수", "상품 목록", "판매 가능한 상품의 대표 이미지, 상품명, 가격 및 판매 상태를 표시한다.", "판매 중·품절 구분"),
    ("USR-03", "운영 필수", "상품 상세", "상품 대표 이미지, 설명, 판매가격 및 구매 가능한 재고 정보를 조회한다.", "대표 이미지 1장 기준"),
    ("USR-04", "필수", "수량 선택", "이용일·이용시간·회차 없이 상품별 구매 수량만 선택한다.", "재고·구매 제한 초과 불가"),
    ("USR-05", "필수", "복수 상품 장바구니", "서로 다른 종류의 굿즈를 한 장바구니에 담아 함께 구매할 수 있다.", "예약 상품 혼합 제외"),
    ("USR-06", "운영 필수", "장바구니 관리", "상품별 수량을 변경하거나 상품을 삭제하고 상품금액·배송비·총액을 확인한다.", "배송비 정책 협의"),
    ("USR-07", "운영 필수", "주문 정보 입력", "주문자명, 연락처, 수령인, 배송지 주소 및 배송 요청사항을 입력한다.", "필수 입력값 협의"),
    ("USR-08", "운영 필수", "재고 재확인", "결제 요청 직전에 최신 재고를 확인하고 부족하면 결제를 중단하여 안내한다.", "장바구니 재고 선점 없음"),
    ("USR-09", "필수", "결제", "기존 회원 로그인과 기존 PG 결제 시스템을 이용하여 카드 결제를 처리한다.", "PG 연동 방식 사전 확인"),
    ("USR-10", "운영 필수", "주문 완료", "결제 성공 후 주문을 생성하고 주문번호와 주문 결과를 표시한다.", "결제 실패 시 주문 미확정"),
    ("USR-11", "운영 필수", "주문 조회", "본인의 주문상품·수량·결제금액·배송지·주문 상태를 조회한다.", "기존 마이페이지 연계"),
    ("USR-12", "필수", "주문 전체 취소", "취소 가능한 출고 전 주문을 주문 단위로 전액 취소한다. 상품별 부분취소는 제공하지 않는다.", "취소 주체·시점 협의"),
]
add_table(doc, ["ID", "구분", "기능", "기능 설명", "비고"], user_rows, [2.0, 2.2, 3.0, 6.4, 2.4], {0, 1})

heading(doc, "5. 관리자 기능")
heading(doc, "5.1 굿즈 판매 현황", 2)
admin_order_rows = [
    ("ADM-01", "필수", "메뉴 추가", "기존 예약 운영 관리자에 굿즈 판매 현황 메뉴를 추가한다.", "별도 관리자 구축 없음"),
    ("ADM-02", "운영 필수", "주문 목록", "주문번호, 주문일, 구매자, 상품·수량, 결제금액 및 주문 상태를 조회한다.", "검색 조건은 협의"),
    ("ADM-03", "운영 필수", "주문 상세", "주문상품, 결제정보, 수령인, 연락처, 배송지 및 배송 요청사항을 조회한다.", "개인정보 권한 적용"),
    ("ADM-04", "운영 필수", "출고 대상 엑셀", "출고 대기 주문의 배송 및 상품 정보를 엑셀 파일로 다운로드한다.", "양식 사전 확정"),
    ("ADM-05", "운영 필수", "출고 완료 일괄 처리", "선택한 출고 대기 주문을 출고 완료 상태로 일괄 변경한다.", "주 1회 출고 기준"),
    ("ADM-06", "필수", "주문 전체 취소", "출고 전 주문을 전액 취소하고 PG 취소 성공 시 전체 취소 상태로 변경한다.", "부분취소 제외"),
    ("ADM-07", "운영 필수", "취소 재고 복원", "전체 취소 완료 시 해당 주문에 포함된 모든 상품의 재고를 복원한다.", "중복 복원 방지"),
    ("ADM-08", "협의", "송장번호 입력", "주문별 택배사와 송장번호를 수동 입력하고 사용자 주문 조회에 표시한다.", "택배사 API 미연동"),
]
add_table(doc, ["ID", "구분", "기능", "기능 설명", "비고"], admin_order_rows, [2.0, 2.2, 3.0, 6.4, 2.4], {0, 1})

heading(doc, "5.2 온라인 스토어 운영", 2)
admin_product_rows = [
    ("PRD-01", "필수", "메뉴 추가", "기존 예약 운영 관리자에 온라인 스토어 운영 메뉴를 추가한다.", "별도 관리자 구축 없음"),
    ("PRD-02", "운영 필수", "상품 등록·수정", "상품명, 간단한 설명, 판매가격 및 상품 정보를 등록하고 수정한다.", "초기 상품 수 협의"),
    ("PRD-03", "운영 필수", "대표 이미지", "상품별 대표 이미지 1장을 등록·교체한다.", "상세 이미지 추가 제외"),
    ("PRD-04", "운영 필수", "재고 관리", "상품별 현재 재고를 입력하고 관리자가 재고를 증감 조정한다.", "변경 사유 기록"),
    ("PRD-05", "운영 필수", "판매 상태", "상품 상태를 판매 중, 품절 또는 숨김으로 변경한다.", "재고 0 자동 품절 협의"),
    ("PRD-06", "운영 필수", "노출 순서", "사용자 상품 목록에서 상품이 표시되는 순서를 설정한다.", "숫자 순서값 기준"),
]
add_table(doc, ["ID", "구분", "기능", "기능 설명", "비고"], admin_product_rows, [2.0, 2.2, 3.0, 6.4, 2.4], {0, 1})

doc.add_page_break()
heading(doc, "6. 주문 결제 재고 처리 기준")
add_table(doc, ["처리 시점", "처리 기준", "실패 또는 예외 처리"], [
    ("장바구니 담기", "재고를 차감하거나 선점하지 않는다.", "결제 시점에 품절될 수 있음을 안내한다."),
    ("결제 직전", "주문상품별 최신 재고와 구매 가능 상태를 다시 확인한다.", "부족한 상품이 있으면 결제를 진행하지 않는다."),
    ("결제 성공", "주문·주문상품·배송정보를 저장하고 상품별 재고를 차감한다.", "주문과 재고 처리는 중복되지 않도록 관리한다."),
    ("결제 실패", "주문을 확정하지 않고 재고를 차감하지 않는다.", "실패 안내 후 장바구니로 복귀할 수 있다."),
    ("전체 취소 성공", "주문의 모든 상품 재고를 복원하고 전체 취소 상태로 변경한다.", "PG 취소 실패 시 재고를 먼저 복원하지 않는다."),
], [4.0, 7.0, 5.0], {0})

heading(doc, "7. 주문 및 출고 상태")
add_table(doc, ["상태", "정의", "사용자 취소", "관리자 처리"], [
    ("결제 완료", "PG 결제가 성공하고 주문이 생성된 상태", "가능", "주문·결제 확인"),
    ("출고 대기", "주간 출고 대상에 포함할 수 있는 상태", "정책에 따라 가능", "엑셀 다운로드·출고 준비"),
    ("출고 완료", "상품 포장 및 택배 접수를 완료한 상태", "불가", "출고 완료 처리·송장 입력"),
    ("전체 취소", "PG 전액 취소와 재고 복원이 완료된 상태", "불가", "조회만 가능"),
], [3.0, 7.0, 3.0, 3.0], {0, 2})
body(doc, "결제 완료에서 출고 대기로 전환하는 시점과 출고 준비 중 취소 잠금 기준은 고객사 운영 정책 확정 후 적용합니다.")

heading(doc, "8. 데이터 구분 원칙")
add_table(doc, ["데이터", "주요 내용", "처리 원칙"], [
    ("상품", "상품명, 설명, 가격, 재고, 판매 상태, 대표 이미지, 노출 순서", "예약 프로그램·회차와 분리"),
    ("주문", "주문번호, 회원, 주문금액, 결제정보, 주문 상태, 주문일시", "한 주문에 여러 상품 포함"),
    ("주문상품", "상품, 상품명·가격 스냅샷, 수량, 상품금액", "주문 당시 정보를 보존"),
    ("배송정보", "수령인, 연락처, 주소, 배송 요청사항, 송장정보", "주문 단위로 저장"),
    ("재고이력", "상품, 증감 수량, 변경 사유, 처리자, 처리일시", "결제·취소·관리자 조정 이력 관리"),
], [3.0, 8.0, 5.0], {0})

heading(doc, "9. MVP 제외 범위")
excluded = [
    ("예약 상품 혼합", "체험 예약권과 굿즈의 혼합 장바구니 및 혼합 결제"),
    ("상품 옵션", "색상, 크기, 포장 유형 등의 옵션 조합"),
    ("할인", "쿠폰, 포인트 및 회원등급 할인"),
    ("취소·반품", "부분취소·부분환불, 교환·반품 자동 접수"),
    ("배송 자동화", "택배사 API, 자동 송장 출력, 배송 자동 추적"),
    ("고객 알림", "주문 상태 알림톡·문자 자동 발송"),
    ("부가 기능", "후기·문의, 찜하기, 추천 상품, 통계 대시보드"),
    ("구매 유형", "비회원 구매, 해외 배송, 복수 배송지 및 분할 배송"),
    ("콘텐츠 제작", "마사회가 별도로 제작하는 인형별 정보제공 페이지 및 신규 상품 이미지 제작"),
]
add_table(doc, ["구분", "제외 내용"], excluded, [4.0, 12.0], {0})

heading(doc, "10. 착수 전 확정 사항")
policy_rows = [
    ("배송비", "무료·정액·조건부 무료 여부와 금액"),
    ("출고", "주간 출고 요일, 주문 마감 시각, 출고 완료 처리 시점"),
    ("취소", "사용자 직접 취소 여부, 취소 가능 상태, 출고 준비 주문 잠금 기준"),
    ("재고", "재고 0 자동 품절 여부, 관리자 재고 조정 권한과 이력 범위"),
    ("구매 제한", "상품별·주문별 최대 수량 및 고액 주문 제한"),
    ("상품 구성", "최초 등록 상품 수, 상품 정보·이미지 제공 일정"),
    ("엑셀", "출고 대상 엑셀 컬럼과 정렬 기준"),
    ("송장", "송장번호 수동 입력 기능의 MVP 포함 여부"),
    ("관리자 권한", "메뉴 접근 대상과 개인정보 조회 권한"),
    ("PG", "상품 주문 결제·전체 취소 API 재사용 가능 여부 및 테스트 환경"),
]
add_table(doc, ["정책", "확정할 내용"], policy_rows, [4.0, 12.0], {0})

heading(doc, "11. 검수 기준")
accept_rows = [
    ("상품 구매", "이용일·시간 선택 없이 상품과 수량만 선택하여 주문할 수 있다."),
    ("복수 상품", "서로 다른 상품을 한 장바구니와 한 주문으로 결제할 수 있다."),
    ("재고", "결제 성공 시에만 재고가 차감되고 전체 취소 완료 시 정확히 복원된다."),
    ("취소", "부분취소 없이 허용된 출고 전 주문의 전체 취소만 처리된다."),
    ("관리자", "기존 관리자에서 상품·재고·주문·출고 기능을 사용할 수 있다."),
    ("범위 분리", "상품 주문 데이터가 예약 프로그램·회차·티켓 데이터와 구분되어 저장된다."),
]
add_table(doc, ["검수 항목", "완료 기준"], accept_rows, [4.0, 12.0], {0})

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("주식회사 위이  |  말마프렌즈 온라인 스토어 기능명세서"), 8, False, "666666")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
